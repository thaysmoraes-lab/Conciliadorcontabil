"""Saídas da conciliação: planilha Excel e relatório HTML autocontido."""
import html
import io
import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

STATUS = {"ok": "Conciliado", "arred": "Conciliado (arredondamento)", "explicada": "Divergência explicada",
          "valor": "Divergente", "erro": "Erro na regra", "sem_origem": "Origem não encontrada", "data": "Data divergente"}
NUM = "#,##0.00;[Red]-#,##0.00"
_A = Font(name="Arial", size=10)
_B = Font(name="Arial", size=10, bold=True, color="FFFFFF")
_T = Font(name="Arial", size=10, bold=True)
_H = PatternFill("solid", fgColor="0F5C4C")


def _data(d):
    return d.strftime("%d/%m/%Y") if isinstance(d, (datetime.date, datetime.datetime)) else (d or "")


def _aba(ws, cab, linhas, larguras, numericas=(), titulo=None):
    r0 = 1
    if titulo:
        ws.cell(1, 1, titulo).font = Font(name="Arial", size=11, bold=True)
        r0 = 3
    for j, h in enumerate(cab, 1):
        c = ws.cell(r0, j, h)
        c.font, c.fill = _B, _H
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, lin in enumerate(linhas, r0 + 1):
        for j, v in enumerate(lin, 1):
            c = ws.cell(i, j, v)
            c.font = _A
            if j in numericas:
                c.number_format = NUM
    for j, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(cab))}{r0 + max(len(linhas), 1)}"
    return r0


def excel(motor, titulo):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    res = motor.resumo_por_lp()
    cab = ["LP", "Seq", "Descrição", "Base", "Linhas CT2", "Chaves", "Valor CT2", "Valor esperado (CT5_VLR01)",
           "Diferença", "Conciliadas", "Explicadas", "Divergentes", "Origem não encontrada",
           "Origem sem lançamento (qtd)", "Origem sem lançamento (valor)"]
    lin = [[r["lp"], r["seq"], r["descricao"], r["base"], r["linhas"], r["chaves"], round(r["ct2"], 2), round(r["esperado"], 2),
            None, r["conciliadas"], r["explicadas"], r["divergentes"], r["sem_origem"], r["sem_lancamento"],
            round(r["valor_sem_lancamento"], 2)] for r in res]
    r0 = _aba(ws, cab, lin, [6, 6, 40, 8, 10, 9, 16, 18, 14, 11, 11, 11, 12, 14, 16], (7, 8, 9, 15), titulo)
    for i in range(r0 + 1, r0 + 1 + len(lin)):
        c = ws.cell(i, 9, f"=G{i}-H{i}")
        c.number_format, c.font = NUM, _A
    t = r0 + 1 + len(lin)
    ws.cell(t, 3, "Total").font = _T
    for j in (5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15):
        L = get_column_letter(j)
        c = ws.cell(t, j, f"=SUM({L}{r0 + 1}:{L}{t - 1})")
        c.font = _T
        if j in (7, 8, 9, 15):
            c.number_format = NUM
    n = t + 2
    ws.cell(n, 1, "Linhas da CT2 fora da conciliação por regra").font = _T
    for k, v in motor.categorias.items():
        n += 1
        ws.cell(n, 1, k).font = _A
        ws.cell(n, 5, v).font = _A
        c = ws.cell(n, 7, motor.categorias_valor.get(k, 0))
        c.font, c.number_format = _A, NUM
    n += 2
    ws.cell(n, 1, "Valor esperado = CT5_VLR01 recalculado sobre as tabelas de origem, com as tabelas relacionadas posicionadas. "
                  "LPs de item são comparados por documento. Tolerância de R$ 0,01 por linha da CT2.").font = Font(name="Arial", size=9, italic=True)

    div = [x for x in motor.resultado if x["status"] not in ("ok",)]
    w2 = wb.create_sheet("Divergências")
    _aba(w2, ["LP", "Seq", "Situação", "Data", "Histórico CT2", "Valor CT2", "Valor esperado", "Diferença", "Explicação / erro"],
         [[x["lp"], x["seq"], STATUS[x["status"]], _data(x["data"]), x["hist"], x["valor"], x["esperado"], None,
           x["explicacao"] or x["erro"] or ""] for x in div], [6, 6, 22, 11, 42, 14, 14, 13, 90], (6, 7, 8))
    for i in range(2, 2 + len(div)):
        c = w2.cell(i, 8, f"=F{i}-G{i}")
        c.number_format, c.font = NUM, _A
    w3 = wb.create_sheet("Origem sem lançamento")
    _aba(w3, ["LP", "Seq", "Tabela base", "Documento", "Fornecedor/cliente", "Data", "Valor pela regra"],
         [[m["lp"], m["seq"], m["base"], m["documento"], m["participante"], _data(m["data"]), m["valor"]] for m in motor.sem_lancamento],
         [6, 6, 11, 14, 16, 11, 16], (7,))
    w4 = wb.create_sheet("Financeiro sem lançamento")
    _aba(w4, ["Classificação", "Tipo doc", "Rec/Pag", "Motivo baixa", "Título", "Cli/For", "Data disp.", "Valor", "Ident. L.A."],
         [[f["classe"], f["tipodoc"], f["recpag"], f["motivo"], f["titulo"], f["clifor"], _data(f["data"]), f["valor"], f["ident_la"]]
          for f in motor.financeiro_sem_lancamento], [62, 9, 8, 11, 26, 10, 11, 14, 10], (8,))
    w5 = wb.create_sheet("CT2 manual e sem LP")
    _aba(w5, ["Tipo", "Rotina", "Data", "Histórico", "Valor"],
         [["Manual" if s["manual"] else "Automático sem LP", s["rotina"], _data(s["data"]), s["hist"], s["valor"]] for s in motor.soltas],
         [20, 11, 11, 60, 14], (5,))
    w6 = wb.create_sheet("Detalhe por chave")
    R = motor.resultado
    _aba(w6, ["LP", "Seq", "Base", "Chave", "Data CT2", "Data origem", "Histórico CT2", "Linhas CT2", "Registros origem",
              "Valor CT2", "Valor esperado", "Diferença", "Situação"],
         [[x["lp"], x["seq"], x["base"], x["chave"], _data(x["data"]), _data(x["data_origem"]), x["hist"], x["linhas"],
           x["registros"], x["valor"], x["esperado"], None, STATUS[x["status"]]] for x in R],
         [6, 6, 8, 40, 11, 11, 42, 9, 10, 14, 14, 12, 24], (10, 11, 12))
    for i in range(2, 2 + len(R)):
        c = w6.cell(i, 12, f"=J{i}-K{i}")
        c.number_format, c.font = NUM, _A
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fm(v):
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def html_relatorio(motor, titulo):
    e = html.escape
    T = motor.totais()
    linhas = []
    for r in motor.resumo_por_lp():
        d = r["ct2"] - r["esperado"]
        sit = []
        if r["divergentes"]:
            sit.append(f'<span class="p bad">{r["divergentes"]} divergente(s)</span>')
        if r["explicadas"]:
            sit.append(f'<span class="p warn">{r["explicadas"]} explicada(s)</span>')
        if r["sem_origem"]:
            sit.append(f'<span class="p bad">{r["sem_origem"]} sem origem</span>')
        if r["sem_lancamento"]:
            sit.append(f'<span class="p info">{r["sem_lancamento"]} sem lançamento</span>')
        linhas.append(f'<tr><td class="m">{r["lp"]}/{r["seq"]}</td><td>{e(r["descricao"])}</td><td class="n">{r["chaves"]}</td>'
                      f'<td class="n">{_fm(r["ct2"])}</td><td class="n">{_fm(r["esperado"])}</td>'
                      f'<td class="n{" red" if abs(d) > 0.009 else ""}">{_fm(d)}</td>'
                      f'<td>{" ".join(sit) or "<span class=\"p ok\">Conciliado</span>"}</td></tr>')
    div = "".join(
        f'<div class="a {"bad" if x["status"] in ("valor", "erro") else "warn"}"><b>{x["lp"]}/{x["seq"]} · {e(x["hist"])}</b> '
        f'({_data(x["data"])})<br><span class="m">CT2 {_fm(x["valor"])} · esperado {_fm(x["esperado"])} · diferença {_fm(x["dif"])}</span>'
        f'<br>{e(x["explicacao"] or x["erro"] or "Sem explicação automática — revisar.")}</div>'
        for x in motor.resultado if x["status"] in ("valor", "erro", "explicada"))
    sl = "".join(f'<tr><td class="m">{m["lp"]}/{m["seq"]}</td><td class="m">{e(m["documento"])}</td><td class="m">{e(m["participante"])}</td>'
                 f'<td>{_data(m["data"])}</td><td class="n">{_fm(m["valor"])}</td></tr>' for m in motor.sem_lancamento)
    fin = {}
    for f in motor.financeiro_sem_lancamento:
        a = fin.setdefault(f["classe"], [0, 0.0])
        a[0] += 1
        a[1] += f["valor"]
    fin_html = "".join(f"<tr><td>{e(k)}</td><td class='n'>{v[0]}</td><td class='n'>{_fm(v[1])}</td></tr>" for k, v in fin.items())
    cat = "".join(f"<tr><td>{e(k)}</td><td class='n'>{v}</td><td class='n'>{_fm(motor.categorias_valor.get(k, 0))}</td></tr>"
                  for k, v in motor.categorias.items())
    pct = round((T["conciliadas"] + T["explicadas"]) / T["chaves"] * 100, 1) if T["chaves"] else 0
    cards = "".join(f'<div class="c"><div class="v">{v}</div><div class="l">{l}</div></div>' for v, l in [
        (f'{T["linhas"]:,}'.replace(",", "."), "linhas da CT2 conciliadas por regra"), (f"{pct}%".replace(".", ","), "chaves conciliadas ou explicadas"),
        (T["explicadas"], "divergências explicadas"), (T["divergentes"], "divergências em aberto"),
        (T["sem_origem"], "chaves sem origem"), (T["sem_lancamento"], "origem sem lançamento")])
    return f"""<!DOCTYPE html><html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{e(titulo)}</title><style>
:root{{--acc:#0f5c4c;--line:#e2e0d8;--mut:#5d6660;--bg:#f6f5f1}}
body{{font:14px/1.5 -apple-system,"Segoe UI",Roboto,Arial,sans-serif;color:#1d2320;background:var(--bg);margin:0;padding:28px}}
main{{max-width:1150px;margin:auto;background:#fff;border:1px solid var(--line);border-radius:10px;padding:26px 30px}}
h1{{font-size:21px;margin:0 0 4px}} h2{{font-size:15px;margin:26px 0 8px}} .sub{{color:var(--mut);margin:0 0 18px}}
.cs{{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:10px}} .c{{background:var(--bg);border-radius:8px;padding:11px 13px}}
.v{{font-size:21px;font-weight:600}} .l{{font-size:12px;color:var(--mut)}}
table{{border-collapse:collapse;width:100%;font-size:12.5px}} th{{text-align:left;color:var(--mut);font-weight:500;border-bottom:1px solid var(--line);padding:6px}}
td{{border-bottom:1px solid var(--line);padding:6px;vertical-align:top}} .n{{text-align:right;font-family:Consolas,monospace}} .m{{font-family:Consolas,monospace}}
.red{{color:#a12f2f}} .p{{font-size:11px;border-radius:5px;padding:1px 7px;white-space:nowrap}} .ok{{background:#e8f2dd;color:#2f6b1f}}
.warn{{background:#fbf0d9;color:#8a5a06}} .bad{{background:#fbe9e9;color:#a12f2f}} .info{{background:#e6eff9;color:#1f4f86}}
.a{{border-radius:8px;padding:10px 12px;margin-bottom:8px;font-size:13px}} .tw{{overflow-x:auto}}
</style></head><body><main><h1>{e(titulo)}</h1>
<p class="sub">Gerado em {datetime.datetime.now():%d/%m/%Y %H:%M} · valor esperado = CT5_VLR01 recalculado sobre as tabelas de origem.</p>
<div class="cs">{cards}</div>
<h2>Resumo por LP e sequência</h2><div class="tw"><table><thead><tr><th>LP/seq</th><th>Descrição</th><th class="n">Chaves</th><th class="n">Valor CT2</th>
<th class="n">Esperado</th><th class="n">Diferença</th><th>Situação</th></tr></thead><tbody>{"".join(linhas)}
<tr><td></td><td><b>Total</b></td><td class="n"><b>{T["chaves"]}</b></td><td class="n"><b>{_fm(T["ct2"])}</b></td><td class="n"><b>{_fm(T["esperado"])}</b></td>
<td class="n"><b>{_fm(T["ct2"] - T["esperado"])}</b></td><td></td></tr></tbody></table></div>
<h2>Divergências</h2>{div or "<p class='sub'>Nenhuma.</p>"}
<h2>Origem com valor pela regra e sem lançamento</h2><div class="tw"><table><thead><tr><th>LP/seq</th><th>Documento</th><th>Forn./cliente</th><th>Data</th><th class="n">Valor</th></tr></thead><tbody>{sl}</tbody></table></div>
<h2>Movimentos financeiros sem lançamento</h2><table><thead><tr><th>Classificação</th><th class="n">Qtd</th><th class="n">Valor</th></tr></thead><tbody>{fin_html}</tbody></table>
<h2>Linhas da CT2 fora da conciliação por regra</h2><table><thead><tr><th>Categoria</th><th class="n">Linhas</th><th class="n">Valor</th></tr></thead><tbody>{cat}</tbody></table>
</main></body></html>"""
